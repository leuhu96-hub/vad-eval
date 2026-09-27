"""Xuất kết quả đánh giá ra Excel."""
import json
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.drawing.image import Image as XLImage
from openpyxl.comments import Comment

R = "results/"; S = json.load(open(R + "summary.json"))
pf = pd.read_csv(R + "per_file.csv", keep_default_na=False); tA = pd.read_csv(R + "tierA.csv"); tB = pd.read_csv(R + "tierB.csv")
thr = pd.read_csv(R + "thr_sweep_dev.csv"); cat = pd.read_csv(R + "category.csv"); abl = pd.read_csv(R + "ablation.csv")
grid = pd.read_csv(R + "grid_dev.csv"); snr = pd.read_csv(R + "snr_auc.csv"); sus = pd.read_csv(R + "suspicious.csv")
lag = pd.read_csv(R + "lag_per_file.csv", keep_default_na=False)

F = "Arial"; HDR = PatternFill("solid", fgColor="1F4E78"); SEC = PatternFill("solid", fgColor="DDEBF7")
thin = Side(style="thin", color="BFBFBF"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical="top"); CE = Alignment(horizontal="center", vertical="top", wrap_text=True)
PCT = "0.0%"; D3 = "0.000"; D4 = "0.0000"

def title(ws, t, sub=None):
    ws["A1"] = t; ws["A1"].font = Font(name=F, size=14, bold=True, color="1F4E78")
    if sub: ws["A2"] = sub; ws["A2"].font = Font(name=F, size=9, italic=True, color="595959")
    ws.sheet_view.showGridLines = False

def table(ws, r0, df, widths=None, fmts=None, center=True):
    fmts = fmts or {}
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(row=r0, column=j, value=c); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = HDR; x.alignment = CE; x.border = BD
        if widths: ws.column_dimensions[x.column_letter].width = widths[j - 1]
    for i, row in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, (np.floating, float)) and np.isnan(v): v = None
            if isinstance(v, (np.integer,)): v = int(v)
            if isinstance(v, (np.floating,)): v = float(v)
            if isinstance(v, (np.bool_,)): v = bool(v)
            x = ws.cell(row=r0 + i, column=j, value=v); x.font = Font(name=F, size=10); x.border = BD
            x.alignment = CE if (center and j > 1) else WR
            if df.columns[j - 1] in fmts: x.number_format = fmts[df.columns[j - 1]]
    return r0 + len(df)

def status_fmt(ws, rng):
    for txt, col in (("Đạt", "C6EFCE"), ("Cảnh báo", "FFEB9C"), ("Không đạt", "FFC7CE"), ("Không áp dụng", "EDEDED")):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=PatternFill("solid", fgColor=col)))

wb = Workbook()
# ======================= Tổng quan =======================
ws = wb.active; ws.title = "Tong quan"
title(ws, "Kết quả đánh giá VAD trên bộ dữ liệu giả định", "Dữ liệu và model đều là mô phỏng để minh họa quy trình. Số liệu không đại diện cho model thật.")
ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 70
rows = [("Bộ dữ liệu", None, None),
 ("Số clip GT / số file model", "150 / 150", "5 category × 30 clip 10s; 1 cặp lệch tên được cài sẵn"),
 ("Cặp hợp lệ sau ghép file", S["n_pairs"], "Music_31 thiếu GT, Gt_babble_30 thiếu model → loại"),
 ("File dùng nguồn âm đã có trong train", S["pct_seen"], "Loại khỏi tập đánh giá chính (rò rỉ)"),
 ("Dev / Test (chia theo nhóm câu nói gốc)", f'{S["n_dev"]} / {S["n_test"]}', "Chỉ gồm file 'unseen'"),
 ("Kết quả chính", None, None),
 ("AUC gộp toàn bộ (cách tính ban đầu)", S["auc_pooled_all"], "Con số 'đẹp' nếu không kiểm tra dữ liệu"),
 ("AUC tầng A trên test (unseen)", float(tA.AUC_micro[0]), f'CI 95% bootstrap theo file: {S["AUC_test_CI"][0]:.3f}–{S["AUC_test_CI"][1]:.3f}'),
 ("Ngưỡng chọn trên dev (DCF-min, qua pipeline A4)", S["thr_used"], f'Youden trên điểm thô dev = {S["thr_youden_raw_dev"]:.2f}'),
 ("F1 strict (test, A4)", float(tB.F1[0]), f'CI 95%: {S["F1_CI"][0]:.3f}–{S["F1_CI"][1]:.3f}'),
 ("F1 collar ±1 bin (test, A4)", float(tB.F1[1]), f'CI 95%: {S["F1nb_CI"][0]:.3f}–{S["F1nb_CI"][1]:.3f}'),
 ("DCF strict (test, A4)", float(tB.DCF[0]), f'CI 95%: {S["DCF_CI"][0]:.3f}–{S["DCF_CI"][1]:.3f}'),
 ("Tỉ lệ lỗi thật (MSC+NDS) / tổng lỗi", S["pct_true_err"], "Phần còn lại là lỗi biên (FEC+OVER)"),
 ("Event-F1 (collar 0.5s) / onset-only", f'{S["event_f1"]:.3f} / {S["event_f1_onset"]:.3f}', None),
]
r = 4
for k, v, n in rows:
    a = ws.cell(row=r, column=1, value=k)
    if v is None and n is None:
        a.font = Font(name=F, bold=True, color="FFFFFF")
        for c in (1, 2, 3): ws.cell(row=r, column=c).fill = HDR
    else:
        a.font = Font(name=F, bold=True); b = ws.cell(row=r, column=2, value=v); b.font = Font(name=F); b.alignment = Alignment(horizontal="center")
        if isinstance(v, float): b.number_format = PCT if v < 1 and k.startswith(("File dùng", "Tỉ lệ")) else D3
        c = ws.cell(row=r, column=3, value=n); c.font = Font(name=F, size=9); c.alignment = WR
    r += 1
worst = cat.sort_values("F1_strict").iloc[0]; a4 = abl.set_index("Cấu_hình")
findings = [
 f'Rò rỉ: AUC trên file dùng nguồn âm đã thấy khi train = {S["auc_seen"]:.4f}, trên file chưa thấy = {S["auc_unseen"]:.4f}. Chênh {S["auc_seen_gap"]:.3f} ≥ 0.01 → chỉ báo cáo trên file chưa thấy.',
 f'Union-find theo cả câu nói và nhiễu tạo một nhóm khổng lồ {S["largest_group_full"]}/{len(pf)} file → không chia được. Phải chia theo câu nói gốc; còn {S["noise_shared_dev_test"]} nguồn nhiễu dùng chung dev/test.',
 f'Offset tâm cửa sổ: đặt score tại start, F1 (ngưỡng 0.5) ở shift 0 chỉ {S["F1_start_at0"]:.3f}, đỉnh dời sang +{S["best_shift_start"]:.2f}s; đặt tại tâm, đỉnh ở 0 với F1 {S["F1_center_at0"]:.3f}.',
 f'Lag từng file phát hiện đúng 3 file cài lỗi lệch 1 bin ({", ".join(S["lag_flag_files"])}), đã sửa. {len(S["lag_review_files"])} file khác có lag 0.2–0.4s, đều SNR thấp/babble: đây là nhiễu ước lượng, không phải lỗi.',
 f'Ngưỡng: Youden điểm thô ({S["thr_youden_raw_dev"]:.2f}) khác ngưỡng tối ưu qua pipeline ({S["thr_used"]:.2f}). Áp Youden cho A4 trên test cho DCF {S["dcf_A4_at_youden"]:.3f} so với {S["dcf_A4_at_best"]:.3f}.',
 f'Lỗi chủ yếu ở biên: F1 tăng từ {tB.F1[0]:.3f} (strict) lên {tB.F1[1]:.3f} khi bỏ bin biên. OVER ({S["err_types"]["OVER"]} bin) là loại lỗi lớn nhất; Δonset TB {S["mean_d_onset"]:+.2f}s, Δoffset {S["mean_d_offset"]:+.2f}s.',
 f'Category yếu nhất: {worst.Category} (F1 {worst.F1_strict:.3f}, FAR {worst.FAR:.1%}). Babble và Music gây false alarm giữa vùng non-speech (NDS).',
 f'Hậu xử lý cải thiện ít: F1 strict A0 {a4.loc["A0","F1_strict"]:.3f} → A4 {a4.loc["A4","F1_strict"]:.3f}. Merge <500ms (A3) làm FAR tăng lên {a4.loc["A3","FAR"]:.1%}; {S["gap_short_filled"]}/{S["gap_short_total"]} khoảng lặng GT ≤1s bị lấp.',
 f'Nhãn: κ bin trong = {S["kappa_interior"]:.2f}, κ bin biên = {S["kappa_boundary"]:.2f}; lệch biên trung vị {S["median_abs_offset_ms"]:.0f} ms. Một phần "lỗi biên" của model là độ mơ hồ của nhãn.',
]
r += 1; ws.cell(row=r, column=1, value="Phát hiện chính").font = Font(name=F, bold=True, color="FFFFFF")
for c in (1, 2, 3): ws.cell(row=r, column=c).fill = HDR
for i, t in enumerate(findings, 1):
    r += 1; ws.cell(row=r, column=1, value=i).font = Font(name=F, bold=True); ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="top")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    c = ws.cell(row=r, column=2, value=t); c.font = Font(name=F, size=10); c.alignment = WR; ws.row_dimensions[r].height = 42

# ======================= Dữ liệu mô phỏng =======================
ws = wb.create_sheet("Du lieu mo phong"); title(ws, "Thiết kế bộ dữ liệu và model giả định")
d = pd.DataFrame([
 ("Asm", 30, "15–30", "Sạch vừa; 3 clip toàn speech", "Không"),
 ("Clean", 30, "20–35", "Sạch", "Không"),
 ("Noise", 30, "−5–10", "SNR thấp → score nhiễu hơn; 1 clip toàn speech", "Ít"),
 ("Music", 30, "0–15", "Nhạc có lời gây điểm cao giả; 5 clip toàn non-speech", "Nhiều"),
 ("Babble", 30, "0–12", "Nền nhiều giọng → điểm nền cao hơn", "Vừa"),
], columns=["Category", "#clip", "SNR (dB)", "Đặc điểm mô phỏng", "Sự kiện giống giọng trong non-speech"])
r = table(ws, 4, d, [12, 8, 10, 50, 22])
r += 2; ws.cell(row=r, column=1, value="Model giả định").font = Font(name=F, bold=True)
for t in ["Cửa sổ 0.96s, hop 0.08s (114 cửa sổ/clip); score = sigmoid(a·(tỉ lệ speech trong cửa sổ, trọng số Hanning − 0.45) + nhiễu AR(1) + sự kiện giả).",
          "Độ dốc a và độ nhiễu phụ thuộc SNR; file có nguồn âm đã thấy khi train được mô phỏng 'nhớ' (nhiễu ×0.5, a ×1.4).",
          "GT tạo ở độ phân giải 10ms, quy về bin 0.5s theo đa số (≥50%). Người gán 2: mỗi biên lệch N(0, 120ms)."]:
    r += 1; ws.cell(row=r, column=1, value="• " + t).font = Font(name=F, size=10)
r += 2
inj = pd.DataFrame([
 ("Model thừa không có GT (Music_31)", "5a", "Có"),
 ("GT không có model (Gt_babble_30)", "5a", "Có"),
 ("3 file lệch timestamp +0.5s (Clean_7, Noise_12, Babble_3)", "5b, 5e", "Có (" + ", ".join(S["lag_flag_files"]) + ")"),
 ("Clip toàn speech / toàn non-speech", "1a", f'Có ({int(pf.degenerate.astype(str).eq("True").sum())} file)'),
 ("~50% clip dùng nguồn âm có trong train", "2a, 2c", f'Có ({S["pct_seen"]:.0%})'),
 ("Nhạc có lời / babble gây false alarm", "4b", "Có (FAR cao ở Babble, Music)"),
 ("Nhãn mơ hồ ở biên (người gán 2)", "3a", f'Có (κ biên {S["kappa_boundary"]:.2f})'),
], columns=["Lỗi cài sẵn", "Mục kiểm tra", "Phát hiện được?"])
table(ws, r, inj, [12, 8, 10, 50, 22], center=False)
for col, w in zip("ABC", [55, 14, 40]): ws.column_dimensions[col].width = max(ws.column_dimensions[col].width, w)

# ======================= Kiểm tra dữ liệu =======================
ws = wb.create_sheet("Kiem tra du lieu"); title(ws, "Checklist kiểm tra chất lượng dữ liệu – kết quả thực hiện", "Trạng thái: Đạt / Cảnh báo / Không đạt / Không áp dụng. Ngưỡng là ngưỡng thực hành.")
def st(cond_bad, cond_warn=False): return "Không đạt" if cond_bad else ("Cảnh báo" if cond_warn else "Đạt")
dpool = S["auc_pooled_all"] - S["auc_pooled_nodeg"]
chk = [
 ("5a", "Ghép file Model↔GT", "Bất kỳ lỗi nào", f'{len(S["problems"])} lỗi: ' + "; ".join(f"{p[0]} ({p[1]})" for p in S["problems"]), st(len(S["problems"]) > 0), "Đã loại 2 file khỏi đánh giá"),
 ("5b", "Thời lượng / số bin", "Lệch > 1 bin", f'{S["n_dur_flag"]} file có start đầu ≠ 0: ' + ", ".join(S["dur_flag_files"]), st(S["n_dur_flag"] > 0), "Trùng với 5e"),
 ("5c", "Định dạng audio (16 kHz mono)", "Khác", "Dữ liệu mô phỏng không có audio", "Không áp dụng", "Với dữ liệu thật: soundfile.info"),
 ("5d", "Hop / cửa sổ", "Sai", f'hop sai: {S["n_hop_bad"]}; cửa sổ sai: {S["n_win_bad"]}; score ngoài [0,1]: {S["n_score_bad"]}', st(S["n_hop_bad"] + S["n_win_bad"] + S["n_score_bad"] > 0), ""),
 ("5e", "Offset tâm cửa sổ / lag từng file", "Đỉnh F1 lệch 0 > ±0.08s; |lag| ≥ 0.4s", f'Đỉnh F1: start {S["best_shift_start"]:+.2f}s, tâm {S["best_shift_center"]:+.2f}s. Lệch 1 bin: ' + ", ".join(S["lag_flag_files"]), st(len(S["lag_flag_files"]) > 0), f'Đã sửa 3 file; {len(S["lag_review_files"])} file "cần xem" là nhiễu ước lượng'),
 ("1a", "File suy biến", "> 5–10%", f'{S["pct_degenerate"]:.1%} số file', st(S["pct_degenerate"] > 0.10, S["pct_degenerate"] > 0.05), "Loại khỏi chỉ số per-file"),
 ("1b", "File cực đoan", "> 30%", f'{S["pct_extreme"]:.1%} số file', st(S["pct_extreme"] > 0.30), ""),
 ("1c", "Độ phức tạp thời gian", "Trung vị transitions ≤ 1", f'Trung vị {S["median_transitions"]:.0f}; bin biên TB {S["mean_boundary_frac"]:.0%}', st(S["median_transitions"] <= 1), "Bin biên chiếm tỉ lệ lớn → phải tách lỗi biên"),
 ("1d", "AUC gộp vs theo file", "Chênh > 0.005; file < 0.9", f'Chênh {dpool:+.4f}; macro {S["auc_macro"]:.3f}±{S["auc_macro_sd"]:.3f}; {S["n_files_auc_lt_09"]} file < 0.9 (tệ nhất {S["auc_worst"]:.3f})', st(abs(dpool) > 0.005, S["n_files_auc_lt_09"] > 0), "Xem timeline các file tệ"),
 ("2a", "Rò rỉ theo metadata", "> 0 file dùng nguồn train", f'{S["pct_seen"]:.0%} file dùng nguồn đã có trong train; nhóm lớn nhất {S["largest_group_full"]} file', st(S["pct_seen"] > 0), "Loại file 'seen'; chia theo câu nói gốc"),
 ("2b", "Nguồn dùng chung dev/test", "> 0", f'{S["groups_crossing"]} nhóm câu nói vắt qua; {S["noise_shared_dev_test"]} nguồn nhiễu dùng chung', st(S["groups_crossing"] > 0, S["noise_shared_dev_test"] > 0), "Nhiễu chung chỉ ảnh hưởng chọn ngưỡng"),
 ("2c", "Tác động của rò rỉ", "ΔAUC ≥ 0.01", f'AUC seen {S["auc_seen"]:.4f} vs unseen {S["auc_unseen"]:.4f} (Δ {S["auc_seen_gap"]:+.4f})', st(S["auc_seen_gap"] >= 0.01), "Báo cáo trên unseen"),
 ("3a", "Đồng thuận nhãn", "κ trong < 0.9; median lệch > 250ms", f'{S["n_ann2_files"]} file; κ trong {S["kappa_interior"]:.2f}, κ biên {S["kappa_boundary"]:.2f}, PABAK {S["pabak"]:.2f}, lệch biên trung vị {S["median_abs_offset_ms"]:.0f}ms', st(S["kappa_interior"] < 0.9 or S["median_abs_offset_ms"] > 250, S["kappa_boundary"] < 0.8), "Báo cáo thêm chế độ collar ±1 bin"),
 ("3b", "Nhãn đáng ngờ", "—", f'{S["n_suspicious"]} bin tự tin trái GT ngoài vùng biên', "Cảnh báo" if S["n_suspicious"] else "Đạt", "Xem sheet 'Nhan dang ngo'; nghe lại"),
 ("4a", "Độ tách điểm số", "< 2% bin mơ hồ và F1 phẳng", f'{S["pct_ambiguous"]:.1%} bin có 0.1<p<0.9 ({S["pct_ambiguous_in_boundary"]:.0%} nằm ở biên); F1 dao động {S["f1_range_0.2_0.9"]:.3f} khi ngưỡng 0.2–0.9', st(S["pct_ambiguous"] < 0.02 and S["f1_range_0.2_0.9"] < 0.02), ""),
 ("4b", "Hard negative", "Không có nhạc có lời/TV/babble", "Có Music (sự kiện giống giọng) và Babble; FAR " + ", ".join(f"{r.Category} {r.FAR:.0%}" for r in cat.itertuples()), "Đạt", "Babble là điểm yếu chính"),
 ("4c", "Phân bố SNR", "> 80% ở trên 15 dB", f'{S["pct_snr_gt15"]:.0%} file > 15 dB; AUC theo bucket: ' + ", ".join(f"{r.snr_bucket}: {r.auc:.3f}" for r in snr.itertuples()), st(S["pct_snr_gt15"] > 0.8), ""),
 ("4d", "Kiểm chứng thực tế", "Chênh lớn so với tổng hợp", "Chưa có bộ test thật", "Không áp dụng", "Cần bổ sung"),
]
df = pd.DataFrame(chk, columns=["Mã", "Kiểm tra", "Ngưỡng cảnh báo", "Kết quả đo", "Trạng thái", "Ghi chú / hành động"])
last = table(ws, 4, df, [6, 24, 24, 62, 13, 34], center=False)
for rr in range(5, last + 1):
    ws.cell(row=rr, column=5).alignment = CE; ws.row_dimensions[rr].height = 44
status_fmt(ws, f"E5:E{last}")
r = last + 2; ws.cell(row=r, column=2, value="Tổng hợp").font = Font(name=F, bold=True)
for i, s in enumerate(["Đạt", "Cảnh báo", "Không đạt", "Không áp dụng"]):
    ws.cell(row=r + 1 + i, column=2, value=s).font = Font(name=F)
    ws.cell(row=r + 1 + i, column=3, value=f'=COUNTIF($E$5:$E${last},B{r+1+i})').font = Font(name=F)
ws.freeze_panes = "C5"

# ======================= Tầng A =======================
ws = wb.create_sheet("Tang A"); title(ws, "Tầng A – điểm thô (trước hậu xử lý), trên test", "AUC_khong_bien: bỏ bin quanh biên GT; AUC_chi_bien: chỉ bin biên. Hàng cuối là file 'seen' để so sánh, không dùng để báo cáo.")
fm = {c: D4 for c in tA.columns if c not in ("Tập", "n_file", "n_bin")}; fm["pct_speech"] = PCT
r = table(ws, 4, tA, [24, 8, 8, 10] + [12] * 8, fm)
ws.conditional_formatting.add(f"E5:E{r}", CellIsRule(operator="lessThan", formula=["0.95"], fill=PatternFill("solid", fgColor="FFC7CE")))
r += 2; ws.cell(row=r, column=1, value="AUC theo bucket SNR (mọi file không suy biến)").font = Font(name=F, bold=True)
table(ws, r + 1, snr, None, {"auc": D4})

# ======================= Chọn ngưỡng =======================
ws = wb.create_sheet("Chon nguong"); title(ws, "Chọn ngưỡng trên dev – quét qua toàn pipeline A4", "DCF = 0.75·MR + 0.25·FAR (công thức). Dòng có DCF nhỏ nhất được đánh dấu.")
g = thr[thr.config == "A4"][["thr", "MR", "FAR", "F1"]].reset_index(drop=True)
r = table(ws, 4, g, [10, 10, 10, 10, 10, 12], {"MR": PCT, "FAR": PCT, "F1": D4})
x = ws.cell(row=4, column=5, value="DCF"); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = HDR; x.alignment = CE
x = ws.cell(row=4, column=6, value="Chọn"); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = HDR; x.alignment = CE
for rr in range(5, r + 1):
    c = ws.cell(row=rr, column=5, value=f"=0.75*B{rr}+0.25*C{rr}"); c.number_format = D4; c.font = Font(name=F, size=10); c.border = BD
    c = ws.cell(row=rr, column=6, value=f'=IF(E{rr}=MIN($E$5:$E${r}),"◄ chọn","")'); c.font = Font(name=F, bold=True, color="C00000")
ws.conditional_formatting.add(f"A5:E{r}", FormulaRule(formula=[f"$E5=MIN($E$5:$E${r})"], fill=PatternFill("solid", fgColor="C6EFCE")))
ws["H4"] = "Ngưỡng chọn (DCF-min)"; ws["I4"] = f'=INDEX(A5:A{r},MATCH(MIN(E5:E{r}),E5:E{r},0))'
ws["H5"] = "Youden trên điểm thô (dev)"; ws["I5"] = round(S["thr_youden_raw_dev"], 3)
ws["H6"] = "DCF test tại ngưỡng Youden"; ws["I6"] = round(S["dcf_A4_at_youden"], 4)
ws["H7"] = "DCF test tại ngưỡng DCF-min"; ws["I7"] = round(S["dcf_A4_at_best"], 4)
ws["H9"] = "Ngưỡng tối ưu từng cấu hình (dev)"
for i, (k, v) in enumerate(S["best_thr_dev"].items()): ws.cell(row=10 + i, column=8, value=k); ws.cell(row=10 + i, column=9, value=v)
for rr in range(4, 15):
    for cc in (8, 9): ws.cell(row=rr, column=cc).font = Font(name=F, bold=(cc == 8 and rr in (4, 9)))
ws.column_dimensions["H"].width = 32; ws.column_dimensions["I"].width = 10; ws.freeze_panes = "A5"

# ======================= Tầng B =======================
ws = wb.create_sheet("Tang B"); title(ws, f'Tầng B – đầu ra A4 trên test (ngưỡng {S["thr_used"]:.2f})', "TP/FN/FP/TN là số bin đo được; các cột tỉ lệ là công thức.")
b = tB[["Chế độ", "TP", "FN", "FP", "TN"]]
r = table(ws, 4, b, [30, 8, 8, 8, 8, 10, 10, 11, 10, 10, 10, 18])
for j, h in enumerate(["MR", "FAR", "Precision", "Recall", "F1", "DCF", "CI 95% F1 (bootstrap file)"], 6):
    x = ws.cell(row=4, column=j, value=h); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = HDR; x.alignment = CE
for rr in range(5, r + 1):
    for j, fml in enumerate([f"=C{rr}/(B{rr}+C{rr})", f"=D{rr}/(D{rr}+E{rr})", f"=B{rr}/(B{rr}+D{rr})", f"=B{rr}/(B{rr}+C{rr})",
                             f"=2*B{rr}/(2*B{rr}+D{rr}+C{rr})", f"=0.75*F{rr}+0.25*G{rr}"], 6):
        c = ws.cell(row=rr, column=j, value=fml); c.number_format = PCT if j != 10 else D4; c.font = Font(name=F, size=10); c.border = BD; c.alignment = CE
ws.cell(row=5, column=12, value=f'{S["F1_CI"][0]:.3f}–{S["F1_CI"][1]:.3f}'); ws.cell(row=6, column=12, value=f'{S["F1nb_CI"][0]:.3f}–{S["F1nb_CI"][1]:.3f}')
r += 2
seg = pd.DataFrame([("FEC* (miss ở bin biên đoạn speech)", S["err_types"]["FEC"], "Biên"), ("MSC (miss giữa đoạn)", S["err_types"]["MSC"], "Thật"),
                    ("OVER* (FA ở bin sát đoạn speech)", S["err_types"]["OVER"], "Biên"), ("NDS (FA giữa vùng non-speech)", S["err_types"]["NDS"], "Thật")],
                   columns=["Loại lỗi", "#bin", "Nhóm"])
r2 = table(ws, r, seg)
ws.cell(row=r2 + 1, column=1, value="% lỗi thật").font = Font(name=F, bold=True)
c = ws.cell(row=r2 + 1, column=2, value=f'=SUMIF(C{r+1}:C{r2},"Thật",B{r+1}:B{r2})/SUM(B{r+1}:B{r2})'); c.number_format = PCT
r = r2 + 3
segm = pd.DataFrame([("Event-F1 (onset+offset, collar 0.5s)", round(S["event_f1"], 4)), ("Event-F1 (onset-only)", round(S["event_f1_onset"], 4)),
                     ("#đoạn GT / #đoạn dự đoán", f'{S["n_gt_seg"]} / {S["n_pred_seg"]}'), ("#đoạn GT bị chẻ (fragmentation)", S["n_fragmented"]),
                     ("#đoạn dự đoán gộp ≥2 đoạn GT (over-merge)", S["n_overmerged"]),
                     ("Δonset trung bình (s)", round(S["mean_d_onset"], 3)), ("Δoffset trung bình (s)", round(S["mean_d_offset"], 3)),
                     ("Khoảng lặng GT ≤1s bị lấp", f'{S["gap_short_filled"]} / {S["gap_short_total"]}')], columns=["Metric theo đoạn", "Giá trị"])
table(ws, r, segm)
ws.cell(row=r + len(segm) + 2, column=1, value="* Định nghĩa đơn giản hóa theo bin: FEC gồm cả cắt ở đầu và cuối đoạn; OVER gồm cả FA ngay trước onset.").font = Font(name=F, size=9, italic=True)

# ======================= Báo cáo category =======================
ws = wb.create_sheet("Bao cao category"); title(ws, f'Báo cáo theo category (test, A4, ngưỡng {S["thr_used"]:.2f})', "DCF, phần F1 mất do biên và % lỗi thật là công thức. CI: bootstrap theo file.")
c2 = cat.rename(columns={"F1_nb": "F1_±1bin"})
r = table(ws, 4, c2, [12, 7, 10, 9, 9, 7, 7, 7, 7, 10, 10, 10, 10, 9, 14, 11],
          {"pct_speech": PCT, "MR": PCT, "FAR": PCT, "F1_strict": D3, "F1_±1bin": D3, "F1_CI_lo": D3, "F1_CI_hi": D3})
for j, h in enumerate(["DCF", "F1 mất do biên", "% lỗi thật"], 14):
    x = ws.cell(row=4, column=j, value=h); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = HDR; x.alignment = CE
for rr in range(5, r + 1):
    for j, fml, nf in ((14, f"=0.75*D{rr}+0.25*E{rr}", PCT), (15, f"=K{rr}-J{rr}", D3), (16, f'=IF(SUM(F{rr}:I{rr})=0,"",(G{rr}+I{rr})/SUM(F{rr}:I{rr}))', "0%")):
        c = ws.cell(row=rr, column=j, value=fml); c.number_format = nf; c.font = Font(name=F, size=10); c.border = BD; c.alignment = CE
t = r + 1; ws.cell(row=t, column=1, value="Macro (TB)").font = Font(name=F, bold=True)
ws.cell(row=t, column=2, value=f"=SUM(B5:B{r})")
for j, L, nf in ((4, "D", PCT), (5, "E", PCT), (10, "J", D3), (11, "K", D3), (14, "N", PCT)):
    c = ws.cell(row=t, column=j, value=f"=AVERAGE({L}5:{L}{r})"); c.number_format = nf; c.fill = SEC; c.font = Font(name=F, bold=True)
for j, L in ((6, "F"), (7, "G"), (8, "H"), (9, "I")): ws.cell(row=t, column=j, value=f"=SUM({L}5:{L}{r})").font = Font(name=F, bold=True)
ws.conditional_formatting.add(f"J5:J{r}", FormulaRule(formula=[f"J5=MIN($J$5:$J${r})"], fill=PatternFill("solid", fgColor="FFC7CE")))
ws.conditional_formatting.add(f"E5:E{r}", FormulaRule(formula=[f"E5=MAX($E$5:$E${r})"], fill=PatternFill("solid", fgColor="FFC7CE")))

# ======================= Ablation =======================
ws = wb.create_sheet("Ablation"); title(ws, "Ablation hậu xử lý (test; ngưỡng chọn riêng cho từng cấu hình trên dev)", "DCF và ΔDCF là công thức.")
a2 = abl.rename(columns={"Cấu_hình": "Cấu hình", "Mô_tả": "Mô tả", "Ngưỡng_dev": "Ngưỡng (dev)", "F1_nb": "F1_±1bin"})
r = table(ws, 4, a2, [10, 30, 11, 9, 9, 10, 10, 7, 7, 7, 7, 10, 12], {"MR": PCT, "FAR": PCT, "F1_strict": D4, "F1_±1bin": D4})
for j, h in enumerate(["DCF", "ΔDCF vs A0"], 12):
    x = ws.cell(row=4, column=j, value=h); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = HDR; x.alignment = CE
for rr in range(5, r + 1):
    c = ws.cell(row=rr, column=12, value=f"=0.75*D{rr}+0.25*E{rr}"); c.number_format = PCT; c.border = BD; c.font = Font(name=F, size=10)
    c = ws.cell(row=rr, column=13, value=f"=L{rr}-$L$5"); c.number_format = "+0.0%;-0.0%;0.0%"; c.border = BD; c.font = Font(name=F, size=10)
r += 2; ws.cell(row=r, column=1, value="Lưới pad × merge trên dev (ngưỡng A4)").font = Font(name=F, bold=True)
gb = grid.sort_values("DCF").reset_index(drop=True)
table(ws, r + 1, gb, None, {"F1": D4, "DCF": D4})
ws.cell(row=r, column=4, value=f'Tốt nhất: pad {int(S["grid_best"]["pre_ms"])}/{int(S["grid_best"]["post_ms"])}ms, merge {int(S["grid_best"]["gap_ms"])}ms (chỉ để tham khảo, không chọn lại trên test)').font = Font(name=F, size=9, italic=True)

# ======================= Theo file =======================
ws = wb.create_sheet("Theo file"); title(ws, "Chỉ số theo từng file")
p2 = pf[["file", "category", "split", "speech_ratio", "transitions", "boundary_frac", "auc_file", "snr_db", "seen", "degenerate", "injected_shift_s"]].merge(lag[["file", "lag_s", "flag"]], on="file", how="left")
p2 = p2.rename(columns={"flag": "cờ_lag"})
r = table(ws, 4, p2, [15, 9, 12, 10, 10, 10, 9, 8, 7, 10, 10, 8, 12], {"speech_ratio": PCT, "boundary_frac": PCT, "auc_file": D3})
ws.conditional_formatting.add(f"G5:G{r}", CellIsRule(operator="lessThan", formula=["0.9"], fill=PatternFill("solid", fgColor="FFC7CE")))
ws.auto_filter.ref = f"A4:M{r}"; ws.freeze_panes = "B5"

# ======================= Nhãn đáng ngờ =======================
ws = wb.create_sheet("Nhan dang ngo"); title(ws, "Bin có score tự tin trái GT (ngoài vùng biên) – cần nghe lại", "Lưu ý: dùng điểm của chính model, chưa phải out-of-fold. Cột 'Kết luận' để bạn điền sau khi nghe.")
s2 = sus.head(40).drop(columns=["nghi_van"]); s2["Kết luận sau khi nghe"] = ""
r = table(ws, 4, s2, [15, 6, 9, 12, 8, 10, 28])
for rr in range(5, r + 1): ws.cell(row=rr, column=7).fill = PatternFill("solid", fgColor="FFF2CC")

# ======================= Biểu đồ =======================
ws = wb.create_sheet("Bieu do"); title(ws, "Biểu đồ")
row = 3
for fn, h in (("fig1_kiem_tra_du_lieu.png", 0.52), ("fig2_tang_A.png", 0.55), ("fig3_tang_B.png", 0.52), ("fig4_timeline.png", 0.55)):
    img = XLImage(R + fn); w, hh = img.width, img.height; img.width, img.height = int(w * h), int(hh * h)
    ws.add_image(img, f"A{row}"); row += int(img.height / 20) + 3

for w in wb.worksheets:
    for rowc in w.iter_rows():
        for c in rowc:
            if c.value is not None and (c.font is None or c.font.name != F):
                c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic, color=c.font.color)
wb.save(R + "Ket_qua_danh_gia_VAD_gia_dinh.xlsx"); print("saved", R + "Ket_qua_danh_gia_VAD_gia_dinh.xlsx")
